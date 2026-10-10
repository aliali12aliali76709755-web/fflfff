"""أدوات الإضافة السريعة للمنتجات والخيارات (نص، Excel، CSV).
- لصق نصي سريع (الاسم | السعر).
- رفع وقراءة ملفات Excel (.xlsx) و CSV مع التحقق والتقرير بالأخطاء.
- إنشاء وتنزيل قالب Excel جاهز.
"""
from __future__ import annotations

import csv
import io
import logging
from typing import Any
import openpyxl

log = logging.getLogger("forge.bulk_import")

HEADERS = ["category", "product_name", "option_name", "price_usd", "cost_provider_usd", "input_type", "input_label"]
AR_HEADERS = ["القسم", "اسم المنتج", "اسم الباقة/الخيار", "سعر البيع ($)", "سعر التكلفة ($)", "نوع المدخل", "عنوان خانة الإدخال"]


class BulkImportManager:
    """إدارة واستيراد المنتجات والخيارات بكميات كبيرة."""

    @staticmethod
    def parse_quick_text(raw_text: str) -> list[dict[str, Any]]:
        """يحلل نصاً سريعاً بصيغة:
        اسم الخيار | السعر
        أو
        القسم | المنتج | الخيار | السعر | التكلفة
        """
        lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]
        results = []
        for line in lines:
            parts = [p.strip() for p in line.split("|")]
            if len(parts) == 2:
                # خيار | سعر
                opt, price_str = parts[0], parts[1]
                try:
                    price = float(price_str.replace("$", "").strip())
                    results.append({
                        "category": "عام",
                        "product_name": "منتج",
                        "option_name": opt,
                        "price_usd": price,
                        "cost_provider_usd": round(price * 0.7, 2),
                    })
                except ValueError:
                    continue
            elif len(parts) >= 4:
                # قسم | منتج | خيار | سعر [| تكلفة]
                cat, prod, opt, price_str = parts[0], parts[1], parts[2], parts[3]
                cost_str = parts[4] if len(parts) > 4 else "0.0"
                try:
                    price = float(price_str.replace("$", "").strip())
                    cost = float(cost_str.replace("$", "").strip())
                    results.append({
                        "category": cat,
                        "product_name": prod,
                        "option_name": opt,
                        "price_usd": price,
                        "cost_provider_usd": cost,
                    })
                except ValueError:
                    continue
        return results

    @classmethod
    def parse_excel_or_csv(cls, file_bytes: bytes, filename: str) -> tuple[list[dict[str, Any]], list[str]]:
        """يقرأ ملف Excel أو CSV ويعيد (الصفوف_الصحيحة, قائمة_الأخطاء)."""
        rows = []
        errors = []

        if filename.endswith(".csv"):
            try:
                content = file_bytes.decode("utf-8-sig")
            except Exception:
                content = file_bytes.decode("latin-1", errors="ignore")
            reader = csv.reader(io.StringIO(content))
            raw_rows = list(reader)
        else:
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active
                raw_rows = list(ws.iter_rows(values_only=True))
            except Exception as e:
                return [], [f"تعذّر فتح ملف Excel: {e}"]

        if not raw_rows or len(raw_rows) < 2:
            return [], ["الملف فارغ أو لا يحتوي بيانات تحت العناوين."]

        # تخطي سطر العناوين
        data_rows = raw_rows[1:]
        for idx, r in enumerate(data_rows, start=2):
            if not r or all(c is None or str(c).strip() == "" for c in r):
                continue
            cat = str(r[0] or "عام").strip()
            prod = str(r[1] or "").strip()
            opt = str(r[2] or "").strip()

            if not prod or not opt:
                errors.append(f"سطر {idx}: اسم المنتج أو الخيار فارغ.")
                continue

            try:
                price = float(str(r[3]).replace("$", "").strip())
                assert price > 0
            except Exception:
                errors.append(f"سطر {idx}: سعر البيع غير صالح ({r[3]}).")
                continue

            try:
                cost = float(str(r[4] if len(r) > 4 and r[4] is not None else "0").replace("$", "").strip())
            except Exception:
                cost = 0.0

            input_type = str(r[5] or "text").strip().lower() if len(r) > 5 and r[5] else "text"
            input_label = str(r[6] or "").strip() if len(r) > 6 and r[6] else ""

            rows.append({
                "category": cat,
                "product_name": prod,
                "option_name": opt,
                "price_usd": price,
                "cost_provider_usd": cost,
                "input_type": input_type,
                "input_label": input_label,
            })

        return rows, errors

    @staticmethod
    def generate_sample_excel() -> bytes:
        """يولد ملف Excel نموذجي (.xlsx) يحتوي أمثلة جاهزة لمنتجات الألعاب والاشتراكات."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products Template"

        # كتابة العناوين
        ws.append(AR_HEADERS)

        # أمثلة نموذجية
        samples = [
            ["شحن ألعاب", "PUBG Mobile", "60 شدة (UC)", 0.99, 0.75, "player_id", "اكتب ID اللاعب (أرقام فقط)"],
            ["شحن ألعاب", "PUBG Mobile", "325 شدة (UC)", 4.80, 3.80, "player_id", "اكتب ID اللاعب (أرقام فقط)"],
            ["شحن ألعاب", "PUBG Mobile", "660 شدة (UC)", 9.50, 7.50, "player_id", "اكتب ID اللاعب (أرقام فقط)"],
            ["اشتراكات وتطبيقات", "Netflix", "اشتراك شهر شاشة خاصة", 3.50, 2.20, "email", "إيميل حسابك للتفعيل"],
            ["اشتراكات وتطبيقات", "YouTube Premium", "اشتراك 3 أشهر رسمي", 5.00, 3.50, "email", "إيميل جيميل للدعوة العائلية"],
            ["بطاقات رقمية", "Apple iTunes", "بطاقة $10 أمريكي", 10.50, 9.80, "text", "لا يتطلب إدخال (تسليم كود فوري)"],
        ]
        for s in samples:
            ws.append(s)

        # ضبط عرض الأعمدة
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
